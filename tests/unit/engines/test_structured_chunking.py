"""Markdown 与页面文档级组合策略。"""

from __future__ import annotations

from pathlib import Path
from typing import cast

import pytest
from docx import Document

from comet_rag.engines.chunkers import (
    ChunkDraft,
    FixedSizeChunker,
    MarkdownSectionStrategy,
    MissingDocumentStructureError,
    PageChunkingStrategy,
)
from comet_rag.engines.chunkers.strategies import DocumentStructureError
from comet_rag.engines.documents import MarkdownDocumentNormalizer
from comet_rag.engines.documents.docx import DocxDocumentExtractor
from comet_rag.ports import DocumentBlock, ExtractedDocument, NormalizedDocument


def _assert_exact_spans(document: NormalizedDocument, drafts: list[ChunkDraft]) -> None:
    for draft in drafts:
        start, end = _required_span(draft)
        assert document.markdown[start:end] == draft.text


def _required_span(draft: ChunkDraft) -> tuple[int, int]:
    if draft.start_char is None or draft.end_char is None:
        raise AssertionError("expected an exact source span")
    return draft.start_char, draft.end_char


def test_markdown_strategy_keeps_section_boundaries_and_heading_paths() -> None:
    document = MarkdownDocumentNormalizer().normalize(
        ExtractedDocument(
            markdown="# 安装\n\n第一部分较长正文。\n\n## Docker\n\n第二部分较长正文。"
        )
    )

    drafts = MarkdownSectionStrategy(FixedSizeChunker(8, 2)).split(document)

    _assert_exact_spans(document, drafts)
    assert [draft.ordinal for draft in drafts] == list(range(len(drafts)))
    assert {
        tuple(cast("list[str]", draft.metadata["heading_path"])) for draft in drafts
    } == {
        ("安装",),
        ("安装", "Docker"),
    }
    for draft in drafts:
        start, end = _required_span(draft)
        containing = [
            block
            for block in document.blocks
            if block.start_char <= start and end <= block.end_char
        ]
        assert len(containing) == 1


def test_heading_facts_are_not_dropped_as_empty_structure() -> None:
    markdown = (
        "# 紧急维护电话为 6118\n\n## 适用范围\n\n夜间值班。\n\n## 紧急授权码为 5931"
    )
    document = MarkdownDocumentNormalizer().normalize(ExtractedDocument(markdown))
    drafts = MarkdownSectionStrategy(FixedSizeChunker(1000)).split(document)
    _assert_exact_spans(document, drafts)
    assert "".join(d.text for d in drafts) == markdown
    assert any("紧急维护电话为 6118" in d.text for d in drafts)
    assert any("紧急授权码为 5931" in d.text for d in drafts)


def test_docx_fixture_preserves_heading_boundaries_and_document_metadata(
    tmp_path: Path,
) -> None:
    path = tmp_path / "structure.docx"
    source = Document()
    source.add_heading("安装", level=1)
    source.add_paragraph("安装说明。")
    source.add_heading("Docker", level=2)
    source.add_paragraph("容器说明。")
    source.save(str(path))
    extracted = DocxDocumentExtractor().extract(
        path,
        filename=path.name,
        media_type=(
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
    )
    document = MarkdownDocumentNormalizer().normalize(extracted)

    drafts = MarkdownSectionStrategy(FixedSizeChunker(100)).split(document)

    _assert_exact_spans(document, drafts)
    assert document.metadata["file_name"] == "structure.docx"
    assert [draft.metadata["heading_path"] for draft in drafts] == [
        ["**安装**"],
        ["**安装**", "**Docker**"],
    ]


def test_page_strategy_never_crosses_pages_and_preserves_page_numbers() -> None:
    markdown = "第一页正文ABCDE第二页正文FGHIJ"
    boundary = len("第一页正文ABCDE")
    document = NormalizedDocument(
        markdown=markdown,
        blocks=(
            DocumentBlock("page", 0, 0, boundary, page_number=1),
            DocumentBlock("page", 1, boundary, len(markdown), page_number=2),
        ),
    )

    drafts = PageChunkingStrategy(FixedSizeChunker(5, 1)).split(document)

    _assert_exact_spans(document, drafts)
    assert {draft.metadata["page_number"] for draft in drafts} == {1, 2}
    for draft in drafts:
        _, end = _required_span(draft)
        page = draft.metadata["page_number"]
        assert page == (1 if end <= boundary else 2)


@pytest.mark.parametrize(
    "strategy",
    [
        MarkdownSectionStrategy(FixedSizeChunker(10)),
        PageChunkingStrategy(FixedSizeChunker(10)),
    ],
)
def test_structure_strategies_refuse_to_guess_missing_facts(strategy: object) -> None:
    with pytest.raises(MissingDocumentStructureError):
        strategy.split(NormalizedDocument(markdown="正文"))  # type: ignore[attr-defined]


def test_page_strategy_requires_page_numbers_to_be_strictly_increasing() -> None:
    document = NormalizedDocument(
        markdown="一二",
        blocks=(
            DocumentBlock("page", 0, 0, 1, page_number=2),
            DocumentBlock("page", 1, 1, 2, page_number=1),
        ),
    )

    with pytest.raises(DocumentStructureError, match="严格递增"):
        PageChunkingStrategy(FixedSizeChunker(10)).split(document)


def test_strategy_rejects_a_chunker_without_source_spans() -> None:
    class SpanlessChunker:
        def split(self, text: str, /) -> list[ChunkDraft]:
            return [ChunkDraft(text=text, ordinal=0)]

    document = NormalizedDocument(
        markdown="正文",
        blocks=(DocumentBlock("section", 0, 0, 2),),
    )

    with pytest.raises(DocumentStructureError, match="字符 span"):
        MarkdownSectionStrategy(SpanlessChunker()).split(document)


@pytest.mark.parametrize(
    "spans", [[], [(3, 9)], [(0, 3), (6, 9)], [(0, 6)], [(0, 6), (3, 9), (2, 5)]]
)
@pytest.mark.parametrize("kind", ["page", "section"])
def test_structured_chunker_cannot_silently_drop_or_reorder_content(
    spans, kind: str
) -> None:
    class FaultyChunker:
        def split(self, text: str, /) -> list[ChunkDraft]:
            return [
                ChunkDraft(text[start:end], i, start, end)
                for i, (start, end) in enumerate(spans)
            ]

    doc = NormalizedDocument(
        "AAABBBCCC",
        blocks=(
            DocumentBlock(kind, 0, 0, 9, page_number=1 if kind == "page" else None),
        ),
    )
    strategy = (
        PageChunkingStrategy(FaultyChunker())
        if kind == "page"
        else MarkdownSectionStrategy(FaultyChunker())
    )
    with pytest.raises(DocumentStructureError):
        strategy.split(doc)


def test_page_blocks_cannot_hide_nonblank_gaps() -> None:
    doc = NormalizedDocument(
        "AAABBBCCC",
        blocks=(
            DocumentBlock("page", 0, 0, 3, page_number=1),
            DocumentBlock("page", 1, 6, 9, page_number=2),
        ),
    )
    with pytest.raises(DocumentStructureError, match="遗漏"):
        PageChunkingStrategy(FixedSizeChunker(9)).split(doc)


@pytest.mark.parametrize(
    "text,spans",
    [
        ("AAABBBCCC", [(0, 6), (3, 9)]),
        ("AAABBBCCC", [(0, 9), (3, 6)]),
        ("  AAA  BBB  ", [(2, 5), (7, 10)]),
        ("  \n ", []),
    ],
)
def test_coverage_validation_allows_overlap_and_only_whitespace_gaps(
    text, spans
) -> None:
    class Chunker:
        def split(self, value: str, /) -> list[ChunkDraft]:
            return [
                ChunkDraft(value[start:end], i, start, end)
                for i, (start, end) in enumerate(spans)
            ]

    doc = NormalizedDocument(
        text, blocks=(DocumentBlock("page", 0, 0, len(text), page_number=1),)
    )
    result = PageChunkingStrategy(Chunker()).split(doc)
    assert [(d.start_char, d.end_char) for d in result] == spans
