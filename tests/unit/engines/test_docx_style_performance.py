from __future__ import annotations

from collections import Counter

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.styles import CT_Styles
from docx.styles.style import CharacterStyle

from comet_rag.engines.documents.docx import DocxDocument, DocxParser


@pytest.mark.parametrize("in_table", [False, True], ids=["paragraph", "table"])
def test_default_style_scans_do_not_grow_with_repeated_runs(
    monkeypatch: pytest.MonkeyPatch, in_table: bool
) -> None:
    """统计上游全表扫描，不以硬件相关的毫秒阈值充当性能守卫。"""
    default_for = CT_Styles.default_for

    def parse_and_count(run_count: int) -> Counter[WD_STYLE_TYPE]:
        document = Document()
        normal = document.styles["Normal"]
        assert isinstance(normal, CharacterStyle)
        normal.font.bold = True
        if in_table:
            table = document.add_table(rows=run_count, cols=1)
            for cell in table.column_cells(0):
                cell.text = "text"
        else:
            paragraph = document.add_paragraph()
            for _ in range(run_count):
                paragraph.add_run("text")
        scans: Counter[WD_STYLE_TYPE] = Counter()

        def count(styles: CT_Styles, style_type: WD_STYLE_TYPE):
            scans[style_type] += 1
            return default_for(styles, style_type)

        with monkeypatch.context() as patch:
            patch.setattr(CT_Styles, "default_for", count)
            parsed = DocxParser().parse(DocxDocument(document, {}))
        if in_table:
            assert parsed.blocks[0]["rows"] == [["**text**"] for _ in range(run_count)]
        else:
            assert parsed.text == "**" + "text" * run_count + "**"
        return scans

    one_run = parse_and_count(1)
    many_runs = parse_and_count(40)
    assert one_run[WD_STYLE_TYPE.CHARACTER] > 0
    assert one_run[WD_STYLE_TYPE.PARAGRAPH] > 0
    assert many_runs == one_run, "重复 run/表格行不应重复扫描同一默认样式"


def test_style_inheritance_work_does_not_grow_with_repeated_runs(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    base_style = CharacterStyle.base_style.fget
    assert base_style is not None

    def parse_and_count(run_count: int) -> int:
        document = Document()
        root = document.styles.add_style("Root character", WD_STYLE_TYPE.CHARACTER)
        assert isinstance(root, CharacterStyle)
        root.font.bold = True
        middle = document.styles.add_style("Middle character", WD_STYLE_TYPE.CHARACTER)
        assert isinstance(middle, CharacterStyle)
        middle.base_style = root
        leaf = document.styles.add_style("Leaf character", WD_STYLE_TYPE.CHARACTER)
        assert isinstance(leaf, CharacterStyle)
        leaf.base_style = middle
        paragraph = document.add_paragraph()
        for _ in range(run_count):
            paragraph.add_run("text").style = leaf
        visits = 0

        def count(style: CharacterStyle):
            nonlocal visits
            visits += 1
            return base_style(style)

        with monkeypatch.context() as patch:
            patch.setattr(CharacterStyle, "base_style", property(count))
            parsed = DocxParser().parse(DocxDocument(document, {}))
        assert parsed.text == "**" + "text" * run_count + "**"
        return visits

    one_run = parse_and_count(1)
    assert one_run > 0
    assert parse_and_count(40) == one_run, "同一样式的继承链不应按 run 重新遍历"
