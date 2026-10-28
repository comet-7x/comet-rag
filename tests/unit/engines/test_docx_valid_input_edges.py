from __future__ import annotations

from typing import cast

import pytest
from docx import Document
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.document import CT_Document
from docx.oxml.ns import qn

from comet_rag.engines.documents.docx import DocxDocument, DocxParser

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def parse(document):
    return DocxParser().parse(DocxDocument(document, {}))


@pytest.mark.parametrize("numbered", [False, True])
def test_default_word_namespace_does_not_require_a_w_prefix(numbered):
    doc = Document()
    properties = (
        '<pPr><numPr><numId n:val="1"/><ilvl n:val="2"/></numPr></pPr>'
        if numbered
        else ""
    )
    root = parse_xml(
        f'<document xmlns="{W}" xmlns:n="{W}"><body><p>{properties}<r><t>Visible text</t></r></p></body></document>'
    )
    doc._element = cast(CT_Document, root)
    assert "w" not in doc.paragraphs[0]._p.nsmap
    assert DocxParser()._get_numId_ilvl(doc.paragraphs[0]) == (
        (1, 2) if numbered else (None, None)
    )
    assert "Visible text" in parse(doc).text


def field(paragraph, instruction: str, *, closed=True):
    for kind in ("begin",):
        element = OxmlElement("w:fldChar")
        element.set(qn("w:fldCharType"), kind)
        paragraph.add_run()._r.append(element)
    command = OxmlElement("w:instrText")
    command.text = instruction
    paragraph.add_run()._r.append(command)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._r.append(separate)
    paragraph.add_run("VISIBLE-FIELD-RESULT")
    if closed:
        end = OxmlElement("w:fldChar")
        end.set(qn("w:fldCharType"), "end")
        paragraph.add_run()._r.append(end)


def test_unclosed_field_preserves_visible_result_and_later_text_without_instruction():
    doc = Document()
    paragraph = doc.add_paragraph("BEFORE ")
    field(paragraph, "REF _Ref123", closed=False)
    paragraph.add_run(" AFTER-FIELD-TEXT")
    text = parse(doc).text
    assert text == "BEFORE VISIBLE-FIELD-RESULT AFTER-FIELD-TEXT"
    assert "REF" not in text


@pytest.mark.parametrize(
    "instruction,caption",
    [
        ("SEQ Figure \\* ARABIC", True),
        ("  seq Table", True),
        ("REF _RefSEQ_ \\r", False),
        ("NOTEREF SEQ_anchor", False),
        ("SEQ_UNKNOWN label", False),
    ],
)
def test_caption_requires_a_seq_instruction_token(instruction, caption):
    doc = Document()
    paragraph = doc.add_paragraph()
    field(paragraph, instruction)
    blocks = parse(doc).blocks
    assert len(blocks) == 1
    assert (blocks[0]["type"] == "caption") is caption
