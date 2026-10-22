from __future__ import annotations

import weakref

import pytest
from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.style import WD_STYLE_TYPE
from docx.opc.constants import RELATIONSHIP_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.styles.style import CharacterStyle, ParagraphStyle
from docx.text.font import Font

from comet_rag.engines.documents.docx import DocxDocument, DocxParser
from comet_rag.engines.documents.docx import parser as parser_module
from comet_rag.engines.documents.docx.formatting import RunFormatResolver

type FontFlags = tuple[bool, bool, bool, bool]


@pytest.fixture
def rendered_formats(monkeypatch: pytest.MonkeyPatch) -> dict[str, FontFlags]:
    """下划线不生成 Markdown 标记，在渲染边界观察四项格式，期望仍手写。"""
    formats: dict[str, FontFlags] = {}
    render = parser_module._apply_fmt_and_url

    def record(text: str, fmt: parser_module._Fmt, url: str | None) -> str:
        formats[text] = (fmt.bold, fmt.italic, fmt.underline, fmt.strikethrough)
        return render(text, fmt, url)

    monkeypatch.setattr(parser_module, "_apply_fmt_and_url", record)
    return formats


def _all_true(font: Font) -> None:
    font.bold = True
    font.italic = True
    font.underline = True
    font.strike = True


def _add_character_style(document: DocumentObject, name: str) -> CharacterStyle:
    style = document.styles.add_style(name, WD_STYLE_TYPE.CHARACTER)
    assert isinstance(style, CharacterStyle)
    return style


def _add_paragraph_style(document: DocumentObject, name: str) -> ParagraphStyle:
    style = document.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(style, ParagraphStyle)
    return style


def _font_style(document: DocumentObject, name: str) -> CharacterStyle:
    style = document.styles[name]
    assert isinstance(style, CharacterStyle)
    return style


def _parse_text(document: DocumentObject, parser: DocxParser | None = None) -> str:
    return (parser or DocxParser()).parse(DocxDocument(document, {})).text


def test_direct_false_overrides_every_inherited_attribute(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    paragraph_style = _add_paragraph_style(document, "Inherited paragraph")
    character_style = _add_character_style(document, "Inherited character")
    _all_true(paragraph_style.font)
    _all_true(character_style.font)

    for attribute in ("bold", "italic", "underline", "strike"):
        paragraph = document.add_paragraph(style=paragraph_style)
        run = paragraph.add_run(attribute)
        run.style = character_style
        setattr(run.font, attribute, False)

    assert (
        _parse_text(document)
        == "~~*bold*~~\n\n~~**italic**~~\n\n~~***underline***~~\n\n***strike***"
    )
    assert rendered_formats == {
        "bold": (False, True, True, True),
        "italic": (True, False, True, True),
        "underline": (True, True, False, True),
        "strike": (True, True, True, False),
    }


def test_character_and_paragraph_chains_preserve_unset_and_explicit_false(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    paragraph_root = _add_paragraph_style(document, "Paragraph root")
    _all_true(paragraph_root.font)
    paragraph_middle = _add_paragraph_style(document, "Paragraph middle")
    paragraph_middle.base_style = paragraph_root
    paragraph_middle.font.italic = False
    paragraph_leaf = _add_paragraph_style(document, "Paragraph leaf")
    paragraph_leaf.base_style = paragraph_middle
    paragraph_leaf.font.underline = False

    character_root = _add_character_style(document, "Character root")
    character_root.font.bold = True
    character_root.font.strike = True
    character_middle = _add_character_style(document, "Character middle")
    character_middle.base_style = character_root
    character_middle.font.italic = True
    character_leaf = _add_character_style(document, "Character leaf")
    character_leaf.base_style = character_middle
    character_leaf.font.underline = True
    paragraph = document.add_paragraph(style=paragraph_leaf)
    paragraph.add_run("character-chain").style = character_leaf

    override_root = _add_character_style(document, "Override root")
    override_root.font.bold = False
    override_middle = _add_character_style(document, "Override middle")
    override_middle.base_style = override_root
    override_middle.font.underline = False
    override_leaf = _add_character_style(document, "Override leaf")
    override_leaf.base_style = override_middle
    override_leaf.font.strike = False
    paragraph = document.add_paragraph(style=paragraph_root)
    paragraph.add_run("character-false").style = override_leaf
    document.add_paragraph("paragraph-chain", style=paragraph_leaf)

    assert (
        _parse_text(document)
        == "~~***character-chain***~~\n\n*character-false*\n\n~~**paragraph-chain**~~"
    )
    assert rendered_formats == {
        "character-chain": (True, True, True, True),
        "character-false": (False, True, False, False),
        "paragraph-chain": (True, False, False, True),
    }


def test_hyperlink_style_is_skipped_by_id_or_name_but_direct_formatting_is_kept(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    _font_style(document, "Normal").font.italic = True
    by_id = _add_character_style(document, "Link display")
    by_id.style_id = "HyperlinkById"
    by_name = _add_character_style(document, "Hyperlink display")
    by_name.style_id = "LinkByName"
    _all_true(by_id.font)
    _all_true(by_name.font)
    paragraph = document.add_paragraph()
    paragraph.add_run("id").style = by_id
    paragraph = document.add_paragraph()
    paragraph.add_run("name").style = by_name

    paragraph = document.add_paragraph()
    run = paragraph.add_run("link")
    run.style = by_id
    run.bold = True
    run.underline = False
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(
        qn("r:id"),
        paragraph.part.relate_to(
            "https://example.test/", RELATIONSHIP_TYPE.HYPERLINK, is_external=True
        ),
    )
    hyperlink.append(run._element)
    paragraph._element.append(hyperlink)

    assert (
        _parse_text(document) == "*id*\n\n*name*\n\n[***link***](https://example.test/)"
    )
    assert rendered_formats == {
        "id": (False, True, False, False),
        "name": (False, True, False, False),
        "link": (True, True, False, False),
    }


def test_field_result_and_table_cell_keep_paragraph_style_inheritance(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    style = _add_paragraph_style(document, "Body emphasis")
    style.font.bold = True
    style.font.italic = True
    paragraph = document.add_paragraph(style=style)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    paragraph.add_run()._element.append(begin)
    instruction = OxmlElement("w:instrText")
    instruction.text = 'HYPERLINK "https://example.test/field"'
    paragraph.add_run()._element.append(instruction)
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    paragraph.add_run()._element.append(separate)
    paragraph.add_run("field").bold = False
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._element.append(end)

    cell_paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    cell_paragraph.style = style
    cell_paragraph.add_run("cell")

    assert _parse_text(document) == (
        "[*field*](https://example.test/field)\n\n| ***cell*** |\n| --- |"
    )
    assert rendered_formats == {
        "field": (False, True, False, False),
        "cell": (True, True, False, False),
    }


def test_missing_and_wrong_type_style_ids_use_the_document_defaults(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    character_default = _font_style(document, "Default Paragraph Font")
    character_default.font.bold = True
    character_default.font.italic = False
    paragraph_default = _font_style(document, "Normal")
    paragraph_default.font.italic = True
    paragraph_default.font.underline = True
    paragraph_default.font.strike = True

    document.add_paragraph("default")
    paragraph = document.add_paragraph()
    paragraph._p.style = "MissingParagraph"
    run = paragraph.add_run("missing")
    run._r.style = "MissingCharacter"
    paragraph = document.add_paragraph()
    paragraph._p.style = "DefaultParagraphFont"
    run = paragraph.add_run("wrong-type")
    run._r.style = "Normal"

    assert (
        _parse_text(document)
        == "~~**default**~~\n\n~~**missing**~~\n\n~~**wrong-type**~~"
    )
    assert rendered_formats == {
        "default": (True, False, True, True),
        "missing": (True, False, True, True),
        "wrong-type": (True, False, True, True),
    }


def test_missing_default_styles_leave_unset_attributes_plain(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    for style in document.styles:
        if qn("w:default") in style.element.attrib:
            del style.element.attrib[qn("w:default")]
    paragraph = document.add_paragraph("plain")
    paragraph._p.style = "MissingParagraph"
    paragraph.runs[0]._r.style = "MissingCharacter"

    assert _parse_text(document) == "plain"
    assert rendered_formats == {"plain": (False, False, False, False)}


def test_last_default_of_each_style_type_wins(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    _font_style(document, "Default Paragraph Font").font.bold = True
    _font_style(document, "Normal").font.bold = True
    character_default = _add_character_style(document, "Last character")
    character_default.element.set(qn("w:default"), "1")
    character_default.font.italic = True
    paragraph_default = _add_paragraph_style(document, "Last paragraph")
    paragraph_default.element.set(qn("w:default"), "1")
    paragraph_default.font.strike = True
    document.add_paragraph("last-default")

    assert _parse_text(document) == "~~*last-default*~~"
    assert rendered_formats == {"last-default": (False, True, False, True)}


def test_missing_base_style_does_not_fall_back_to_the_character_default(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    _font_style(document, "Default Paragraph Font").font.bold = True
    _font_style(document, "Normal").font.strike = True
    orphan = _add_character_style(document, "Orphan character")
    missing_base = OxmlElement("w:basedOn")
    missing_base.set(qn("w:val"), "MissingBase")
    orphan.element.append(missing_base)
    orphan.font.italic = True
    document.add_paragraph().add_run("orphan").style = orphan

    assert _parse_text(document) == "~~*orphan*~~"
    assert rendered_formats == {"orphan": (False, True, False, True)}


def test_doc_defaults_are_not_part_of_the_existing_style_resolution(
    rendered_formats: dict[str, FontFlags],
) -> None:
    """特征化既有边界：此次缓存优化不顺便加入 docDefaults 的新语义。"""
    document = Document()
    defaults = document.styles.element.find(qn("w:docDefaults"))
    assert defaults is not None
    properties = defaults.find(f"{qn('w:rPrDefault')}/{qn('w:rPr')}")
    assert properties is not None
    for tag in ("w:b", "w:i", "w:u", "w:strike"):
        property_element = OxmlElement(tag)
        property_element.set(qn("w:val"), "single" if tag == "w:u" else "1")
        properties.append(property_element)
    document.add_paragraph("doc-defaults")

    assert _parse_text(document) == "doc-defaults"
    assert rendered_formats == {"doc-defaults": (False, False, False, False)}


def test_direct_formatting_does_not_read_a_shadowed_invalid_style_attribute(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    style = _add_character_style(document, "Invalid inherited underline")
    style.font.bold = True
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "invalid-underline")
    properties = style.element.find(qn("w:rPr"))
    assert properties is not None
    properties.append(underline)
    run = document.add_paragraph().add_run("short-circuit")
    run.style = style
    run.underline = False

    assert _parse_text(document) == "**short-circuit**"
    assert rendered_formats == {"short-circuit": (True, False, False, False)}


def test_emphasis_mark_still_overrides_direct_underline_false(
    rendered_formats: dict[str, FontFlags],
) -> None:
    document = Document()
    run = document.add_paragraph().add_run("emphasis")
    run.underline = False
    emphasis = OxmlElement("w:em")
    emphasis.set(qn("w:val"), "dot")
    run._element.get_or_add_rPr().append(emphasis)

    assert _parse_text(document) == "emphasis"
    assert rendered_formats == {"emphasis": (False, False, True, False)}


def test_reusing_parser_does_not_share_default_styles_between_documents() -> None:
    parser = DocxParser()
    first = Document()
    _font_style(first, "Default Paragraph Font").font.bold = True
    first.add_paragraph("first")
    second = Document()
    _font_style(second, "Default Paragraph Font").font.italic = True
    second.add_paragraph("second")

    assert _parse_text(first, parser) == "**first**"
    assert _parse_text(second, parser) == "*second*"
    assert _parse_text(first, parser) == "**first**"


def test_reparsing_document_observes_changed_default_styles() -> None:
    parser = DocxParser()
    document = Document()
    character_default = _font_style(document, "Default Paragraph Font")
    character_default.font.bold = True
    document.add_paragraph("mutable")

    assert _parse_text(document, parser) == "**mutable**"
    character_default.font.bold = False
    _font_style(document, "Normal").font.bold = True
    assert _parse_text(document, parser) == "mutable"
    character_default.element.attrib.pop(qn("w:default"))
    assert _parse_text(document, parser) == "**mutable**"


def test_failed_parse_does_not_poison_the_next_document(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    parser = DocxParser()
    first = Document()
    _font_style(first, "Default Paragraph Font").font.bold = True
    first.add_paragraph("discarded")
    failure = RuntimeError("injected after body parsing")

    def fail() -> None:
        raise failure

    with monkeypatch.context() as patch:
        patch.setattr(parser, "_add_headers_footers", fail)
        with pytest.raises(RuntimeError) as raised:
            _parse_text(first, parser)
    assert raised.value is failure

    second = Document()
    _font_style(second, "Default Paragraph Font").font.italic = True
    second.add_paragraph("recovered")
    assert _parse_text(second, parser) == "*recovered*"


@pytest.mark.parametrize("fail_after_body", [False, True], ids=["success", "failure"])
def test_parse_does_not_retain_style_resolver_after_completion(
    monkeypatch: pytest.MonkeyPatch,
    fail_after_body: bool,
) -> None:
    """缓存只属于一次 parse；用弱引用观察释放，不依赖缓存字典的形状。"""
    references: list[weakref.ReferenceType[RunFormatResolver]] = []

    class ObservedResolver(RunFormatResolver):
        def __init__(self, document: DocumentObject) -> None:
            super().__init__(document)
            references.append(weakref.ref(self))

    monkeypatch.setattr(parser_module, "RunFormatResolver", ObservedResolver)
    document = Document()
    document.add_paragraph("cached")
    parser = DocxParser()

    def fail() -> None:
        raise RuntimeError("injected after body parsing")

    if fail_after_body:
        monkeypatch.setattr(parser, "_add_headers_footers", fail)
        with pytest.raises(RuntimeError, match="injected after body parsing"):
            _parse_text(document, parser)
    else:
        assert _parse_text(document, parser) == "cached"

    assert len(references) == 1
    assert references[0]() is None
