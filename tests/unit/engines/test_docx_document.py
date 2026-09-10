from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from comet_rag.engines.document import DocxDocumentExtractor
from comet_rag.ports import DocumentExtractorPort, ExtractedDocument
from tests.contracts.document_extractor import DocumentExtractorContract

DOCX_MEDIA_TYPE = (
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
)


class TestDocxDocumentExtractor(DocumentExtractorContract):
    @pytest.fixture
    def document_path(self, tmp_path: Path) -> Path:
        path = tmp_path / "sample.docx"
        document = Document()
        document.add_heading("标题", level=1)
        document.add_paragraph("正文")
        document.save(path)
        return path

    @pytest.fixture
    def filename(self) -> str:
        return "sample.docx"

    @pytest.fixture
    def media_type(self) -> str:
        return DOCX_MEDIA_TYPE

    @pytest.fixture
    def extractor(self) -> DocxDocumentExtractor:
        return DocxDocumentExtractor()

    @pytest.fixture
    def expected_document(self, document_path: Path) -> ExtractedDocument:
        return ExtractedDocument(
            markdown="# **标题**\n\n正文",
            metadata={
                "file_name": "sample.docx",
                "file_type": "docx",
                "file_size": document_path.stat().st_size,
                "media_type": DOCX_MEDIA_TYPE,
            },
        )


def test_docx_extractor_satisfies_port() -> None:
    assert isinstance(DocxDocumentExtractor(), DocumentExtractorPort)
