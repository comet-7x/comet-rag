from __future__ import annotations

import asyncio
import hashlib
import io
import threading
from collections.abc import Callable
from dataclasses import FrozenInstanceError
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.styles.style import CharacterStyle, ParagraphStyle
from lxml import etree

from comet_rag.documents.codec import dumps, loads
from comet_rag.documents.parsers.archive import (
    ArchiveLimits,
    ArchiveResourceLimitExceeded,
)
from comet_rag.documents.parsers.docx import AsyncDocxParser
from comet_rag.documents.parsers.text import AsyncTextParser
from comet_rag.documents.projection import project_document
from comet_rag.documents.schemas import (
    Block,
    InlineCode,
    InlineFormula,
    ParseResult,
    TextRun,
)
from comet_rag.sources import AsyncLocalSourceResolver, LocalResource, LocalSource
from tests.fixtures.docx.build import (
    build_basic,
    build_equations,
    build_grid_gaps_table,
    build_header_footer,
    build_image,
    build_merged_table,
    build_table,
)

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def _resolve(path: Path, *, max_bytes: int = 256 * 1024 * 1024) -> LocalResource:
    async def resolve() -> LocalResource:
        async with AsyncLocalSourceResolver(max_bytes=max_bytes).resolve(
            LocalSource(path)
        ) as resource:
            return resource

    return asyncio.run(resolve())


def _parse(
    path: Path,
    output_dir: Path | None = None,
    *,
    archive_limits: ArchiveLimits | None = None,
    max_bytes: int = 256 * 1024 * 1024,
    max_table_cells: int = 10_000,
) -> ParseResult:
    async def parse() -> ParseResult:
        async with AsyncLocalSourceResolver().resolve(
            LocalSource(path, source_id="fixture")
        ) as resource:
            return await AsyncDocxParser(
                archive_limits=archive_limits,
                max_bytes=max_bytes,
                max_table_cells=max_table_cells,
            ).parse(resource, output_dir=output_dir or path.parent / "assets")

    return asyncio.run(parse())


def _text(block: Block) -> str:
    if block.code is not None:
        return block.code
    return (
        "".join(
            item.text
            for item in block.rich_text.inlines
            if isinstance(item, (TextRun, InlineCode))
        )
        if block.rich_text
        else ""
    )


def _rewrite(path: Path, mutate: Callable[[dict[str, bytes]], None]) -> None:
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    mutate(parts)
    with ZipFile(path, "w", compression=ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)


def _mutate_document(path: Path, mutate: Callable[[etree._Element], None]) -> None:
    def update(parts: dict[str, bytes]) -> None:
        root = etree.fromstring(parts["word/document.xml"])
        mutate(root)
        parts["word/document.xml"] = etree.tostring(root)

    _rewrite(path, update)


def test_local_source_borrows_and_hashes_actual_bytes(tmp_path: Path) -> None:
    path = tmp_path / "input.txt"
    data = b"  first\r\nsecond\rthird\n"
    path.write_bytes(data)

    async def resolve() -> tuple[LocalResource, LocalResource]:
        resolver = AsyncLocalSourceResolver()
        async with resolver.resolve(LocalSource(str(path))) as first:
            assert first.path == path.resolve()
            assert first.source.content_hash == hashlib.sha256(data).hexdigest()
            assert first.size_bytes == len(data)
            assert first.original_name == "input.txt"
            assert first.media_type == "text/plain"
            assert first.source.locator == path.as_uri()
            async with resolver.resolve(LocalSource(path, "logical")) as second:
                assert second.source.source_id == "logical"
                assert second.source.content_hash == first.source.content_hash
            assert path.read_bytes() == data
        return first, second

    first, second = asyncio.run(resolve())
    assert first.source.source_id != second.source.source_id
    assert path.read_bytes() == data
    with pytest.raises(FrozenInstanceError):
        first.source.source_id = "changed"  # type: ignore[misc]
    assert _resolve(path).source.source_id == first.source.source_id


@pytest.mark.parametrize("kind", ["missing", "directory", "budget"])
def test_local_source_rejects_unusable_inputs(tmp_path: Path, kind: str) -> None:
    path = tmp_path / "input.txt"
    if kind == "directory":
        path.mkdir()
    elif kind == "budget":
        path.write_bytes(b"12345")
    with pytest.raises((FileNotFoundError, ValueError)):
        _resolve(path, max_bytes=4)


@pytest.mark.parametrize(
    "data", [b"  first\r\nsecond\rthird\n", b"", "﻿\t中 文\r\n".encode()]
)
def test_text_parser_preserves_decoded_bytes(tmp_path: Path, data: bytes) -> None:
    path = tmp_path / "text.md"
    path.write_bytes(data)
    resource = _resolve(path)
    result = asyncio.run(
        AsyncTextParser().parse(resource, output_dir=tmp_path / "unused")
    )
    assert project_document(result.document).text == data.decode("utf-8")
    assert result.document.source is resource.source
    assert result.report.status == "complete"
    assert not result.document.pages
    assert not (tmp_path / "unused").exists()
    assert loads(dumps(result)) == result


def test_text_parser_strict_encoding_budget_and_changed_input(tmp_path: Path) -> None:
    path = tmp_path / "text.txt"
    path.write_bytes(b"caf\xe9\r\n")
    resource = _resolve(path)
    with pytest.raises(UnicodeDecodeError):
        asyncio.run(AsyncTextParser().parse(resource, output_dir=tmp_path))
    result = asyncio.run(
        AsyncTextParser(encoding="latin-1").parse(resource, output_dir=tmp_path)
    )
    assert project_document(result.document).text == "café\r\n"
    with pytest.raises(ValueError, match="exceeds"):
        asyncio.run(AsyncTextParser(max_bytes=2).parse(resource, output_dir=tmp_path))
    path.write_bytes(b"123456")
    with pytest.raises(ValueError, match="content_hash"):
        asyncio.run(AsyncTextParser().parse(resource, output_dir=tmp_path))


def test_docx_basic_runs_headings_lists_and_order(tmp_path: Path) -> None:
    result = _parse(build_basic(tmp_path / "basic.docx"))
    blocks = result.document.blocks
    assert [_text(block) for block in blocks] == [
        "一级标题",
        "这是第一段正文，用于验证纯文本提取。",
        "加粗与斜体混排。",
        "二级标题",
        "二级标题下的正文。",
        "三级标题",
        "第一项",
        "第二项",
        "嵌套项",
        "有序第一项",
        "有序第二项",
    ]
    assert [
        (block.level, _text(block)) for block in blocks if block.kind == "heading"
    ] == [(1, "一级标题"), (2, "二级标题"), (3, "三级标题")]
    lists = [block for block in blocks if block.kind == "list_item"]
    assert [_text(block) for block in lists] == [
        "第一项",
        "第二项",
        "嵌套项",
        "有序第一项",
        "有序第二项",
    ]
    assert [block.marker for block in lists[-2:]] == ["1.", "2."]
    assert lists[0].layout["numbering_format"] == "bullet"
    assert lists[2].layout["indent_twips"] == {"left": "720", "hanging": "360"}
    rich = blocks[2].rich_text
    assert rich is not None
    runs = [item for item in rich.inlines if isinstance(item, TextRun)]
    assert [run.text for run in runs] == ["加粗", "与", "斜体", "混排。"]
    assert runs[0].style and runs[0].style.bold is True
    assert runs[2].style and runs[2].style.italic is True
    assert all(
        block.locations[0].element_id.startswith("word/document.xml:")
        for block in blocks
        if block.locations[0].element_id
    )
    assert result.document.pages == ()
    assert result.report.status == "complete"


def test_native_tables_preserve_empty_repeated_and_merged_cells(tmp_path: Path) -> None:
    simple = _parse(build_table(tmp_path / "table.docx"))
    table = next(block.table for block in simple.document.blocks if block.table)
    assert table.rows == table.columns == 3
    assert table.cells is not None and len(table.cells) == 9
    by_id = {block.block_id: block for block in simple.document.blocks}
    assert [
        [_text(by_id[member]) for member in cell.block_ids] for cell in table.cells
    ] == [
        [text] for text in ["列A", "列B", "列C", "值1", "值2", "", "值4", "值5", "值6"]
    ]
    assert project_document(simple.document).text.endswith("表格后的正文。")
    result = _parse(build_merged_table(tmp_path / "merged.docx"))
    first, second = [block.table for block in result.document.blocks if block.table]
    assert first.cells is not None and second.cells is not None
    assert [
        (cell.row, cell.column, cell.row_span, cell.column_span)
        for cell in first.cells
        if cell.row_span > 1 or cell.column_span > 1
    ] == [(0, 0, 1, 2), (3, 0, 2, 1), (5, 0, 1, 3)]
    assert [
        (cell.row, cell.column, cell.column_span)
        for cell in second.cells
        if cell.column_span > 1
    ] == [(0, 0, 2), (0, 3, 2), (1, 4, 2)]
    assert project_document(result.document).text.count("Q1") == 2
    assert project_document(result.document).text.count("纵向") == 1
    assert result.report.status == "complete"
    assert loads(dumps(result)) == result


def test_table_grid_gaps_remain_absent_cells(tmp_path: Path) -> None:
    result = _parse(build_grid_gaps_table(tmp_path / "gaps.docx"))
    first, second = [block.table for block in result.document.blocks if block.table]
    assert first.cells is not None and second.cells is not None
    assert [(cell.row, cell.column) for cell in first.cells] == [
        (0, 0),
        (0, 1),
        (0, 2),
        (1, 1),
        (1, 2),
        (2, 0),
        (2, 1),
        (3, 1),
    ]
    assert [
        (cell.row, cell.column, cell.column_span)
        for cell in second.cells
        if cell.row == 1
    ] == [(1, 1, 2), (1, 3, 1)]
    assert "\ty\tz\np\tq\t\n\tm\t" in project_document(result.document).text
    assert result.report.status == "complete"


def test_docx_images_materialize_exact_original_bytes_with_stable_identity(
    tmp_path: Path,
) -> None:
    path = build_image(tmp_path / "image.docx")
    first = _parse(path, tmp_path / "output_one")
    second = _parse(path, tmp_path / "output_two")
    (asset,) = first.document.assets
    with ZipFile(path) as archive:
        original = archive.read("word/media/image1.png")
    assert isinstance(asset.media.path, Path)
    assert asset.media.path.read_bytes() == original
    assert asset.content_hash == hashlib.sha256(original).hexdigest()
    assert asset.media.path.stem == asset.content_hash
    assert asset.media.path.parent == tmp_path / "output_one"
    assert first.document.document_id == second.document.document_id
    assert first.report.status == "complete"
    assert [block.kind for block in first.document.blocks] == [
        "heading",
        "paragraph",
        "image",
        "paragraph",
    ]
    assert first.document.blocks[2].layout["width_emu"] == 914400
    path.unlink()
    assert asset.media.path.read_bytes() == original


def test_omml_is_native_and_has_explicit_limitation(tmp_path: Path) -> None:
    result = _parse(build_equations(tmp_path / "math.docx"))
    formulas = [
        inline.formula
        for block in result.document.blocks
        if block.rich_text
        for inline in block.rich_text.inlines
        if isinstance(inline, InlineFormula)
    ]
    assert len(formulas) == 4
    assert all(formula.format == "omml" for formula in formulas)
    roots = [
        etree.fromstring(formula.expression.encode())
        for formula in formulas
        if formula.expression
    ]
    math_ns = "http://schemas.openxmlformats.org/officeDocument/2006/math"
    assert (
        roots[0].findtext(
            f"{{{math_ns}}}f/{{{math_ns}}}num/{{{math_ns}}}r/{{{math_ns}}}t"
        )
        == "a"
    )
    assert (
        roots[0].findtext(
            f"{{{math_ns}}}f/{{{math_ns}}}den/{{{math_ns}}}r/{{{math_ns}}}t"
        )
        == "b"
    )
    assert result.report.status == "partial"
    assert [issue.code for issue in result.report.issues] == ["omml_preserved"] * 4
    assert all(issue.locations for issue in result.report.issues)


def test_header_and_footer_are_native_parts_without_fake_pages(tmp_path: Path) -> None:
    result = _parse(build_header_footer(tmp_path / "hf.docx"))
    assert [(_text(block), block.role) for block in result.document.blocks] == [
        ("页眉页脚样本", None),
        ("正文内容。", None),
        ("这是页眉", "header"),
        ("这是页脚", "footer"),
    ]
    assert result.document.pages == ()
    assert result.report.status == "complete"


def test_nested_tables_links_styles_code_and_inline_image_order(tmp_path: Path) -> None:
    doc = Document()
    base = doc.styles.add_style("BaseCustom", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(base, ParagraphStyle)
    base.font.bold = True
    derived = doc.styles.add_style("DerivedCustom", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(derived, ParagraphStyle)
    derived.base_style = base
    paragraph = doc.add_paragraph(style=derived)
    run = paragraph.add_run("  leading\t")
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x12, 0x34, 0xAB)
    run.font.subscript = True
    run.font.underline = True
    run.bold = False
    hyperlink = OxmlElement("w:hyperlink")
    rid = doc.part.relate_to(
        "https://example.org/a", f"{_R}/hyperlink", is_external=True
    )
    hyperlink.set(qn("r:id"), rid)
    link_run = OxmlElement("w:r")
    link_text = OxmlElement("w:t")
    link_text.text = "link"
    link_run.append(link_text)
    hyperlink.append(link_run)
    paragraph._p.append(hyperlink)
    image_source = build_image(tmp_path / "source.docx")
    with ZipFile(image_source) as archive:
        png = archive.read("word/media/image1.png")
    paragraph.add_run().add_picture(io.BytesIO(png))
    paragraph.add_run(" tail  ")
    code = doc.styles.add_style("Code", WD_STYLE_TYPE.PARAGRAPH)
    assert isinstance(code, ParagraphStyle)
    doc.add_paragraph("  a = 1\n\tb = 2  ", style=code)
    char = doc.styles.add_style("CodeChar", WD_STYLE_TYPE.CHARACTER)
    assert isinstance(char, CharacterStyle)
    doc.add_paragraph().add_run("inline_code").style = char
    outer = doc.add_table(rows=1, cols=1)
    outer.cell(0, 0).text = "outer"
    outer.cell(0, 0).add_table(rows=1, cols=1).cell(0, 0).text = "inner"
    path = tmp_path / "rich.docx"
    doc.save(str(path))
    result = _parse(path)
    assert result.report.status == "complete"
    blocks = result.document.blocks
    assert [block.kind for block in blocks[:5]] == [
        "paragraph",
        "image",
        "paragraph",
        "code",
        "paragraph",
    ]
    assert _text(blocks[0]) == "  leading\tlink"
    assert _text(blocks[2]) == " tail  "
    assert blocks[0].rich_text is not None
    styled = blocks[0].rich_text.inlines[0]
    assert isinstance(styled, TextRun) and styled.style is not None
    assert styled.style.bold is False
    assert styled.style.size_pt == 13
    assert styled.style.color == "#1234AB"
    assert styled.style.script == "subscript"
    assert styled.style.underline is True
    linked = blocks[0].rich_text.inlines[-1]
    assert isinstance(linked, TextRun) and linked.link == "https://example.org/a"
    assert _text(blocks[3]) == "  a = 1\n\tb = 2  "
    assert blocks[4].rich_text and isinstance(
        blocks[4].rich_text.inlines[0], InlineCode
    )
    assert "outer\ninner" in project_document(result.document).text
    assert len([block for block in blocks if block.table]) == 2
    assert loads(dumps(result)) == result


def test_unknown_native_content_is_partial_and_retains_known_descendants(
    tmp_path: Path,
) -> None:
    path = build_table(tmp_path / "unknown.docx")

    def mutate(root: etree._Element) -> None:
        paragraph = root.find(f"{{{_W}}}body/{{{_W}}}p")
        assert paragraph is not None
        unknown = etree.SubElement(paragraph, "{urn:future}meaningful")
        run = etree.SubElement(unknown, f"{{{_W}}}r")
        etree.SubElement(run, f"{{{_W}}}t").text = "retained future text"
        etree.SubElement(paragraph, f"{{{_W}}}altChunk")

    _mutate_document(path, mutate)
    result = _parse(path)
    assert result.report.status == "partial"
    assert "retained future text" in project_document(result.document).text
    assert [issue.code for issue in result.report.issues] == [
        "unsupported_element",
        "unsupported_element",
    ]
    assert all(issue.locations for issue in result.report.issues)


@pytest.mark.parametrize(
    "mutation",
    [
        "missing_image",
        "missing_reference",
        "invalid_xml",
        "unsafe_member",
        "missing_document",
        "doctype",
        "grid_span",
        "grid_gap",
        "vertical_merge",
        "missing_grid",
    ],
)
def test_invalid_docx_resources_fail(tmp_path: Path, mutation: str) -> None:
    path = (
        build_image(tmp_path / "invalid.docx")
        if mutation == "missing_image"
        else build_table(tmp_path / "invalid.docx")
    )

    def mutate(parts: dict[str, bytes]) -> None:
        if mutation == "missing_image":
            del parts["word/media/image1.png"]
        elif mutation == "missing_reference":
            root = etree.fromstring(parts["word/document.xml"])
            body = root.find(f"{{{_W}}}body")
            assert body is not None
            etree.SubElement(body, f"{{{_W}}}altChunk").set(f"{{{_R}}}id", "missing")
            parts["word/document.xml"] = etree.tostring(root)
        elif mutation == "missing_grid":
            root = etree.fromstring(parts["word/document.xml"])
            grid = root.find(f"{{{_W}}}body/{{{_W}}}tbl/{{{_W}}}tblGrid")
            assert grid is not None
            owner = grid.getparent()
            assert owner is not None
            owner.remove(grid)
            parts["word/document.xml"] = etree.tostring(root)
        elif mutation == "invalid_xml":
            parts["word/document.xml"] = b"<not-closed"
        elif mutation == "unsafe_member":
            parts["../escape.bin"] = b"escape"
        elif mutation == "missing_document":
            del parts["word/document.xml"]
        elif mutation == "doctype":
            parts["word/document.xml"] = parts["word/document.xml"].replace(
                b"<w:document",
                b'<!DOCTYPE w:document [<!ENTITY value "injected">]><w:document',
                1,
            )
        else:
            root = etree.fromstring(parts["word/document.xml"])
            cell = root.find(f"{{{_W}}}body/{{{_W}}}tbl/{{{_W}}}tr/{{{_W}}}tc")
            assert cell is not None
            properties = cell.find(f"{{{_W}}}tcPr")
            assert properties is not None
            if mutation == "grid_span":
                etree.SubElement(properties, f"{{{_W}}}gridSpan").set(
                    f"{{{_W}}}val", "10000000"
                )
            elif mutation == "grid_gap":
                row = cell.getparent()
                assert row is not None
                etree.SubElement(
                    etree.SubElement(row, f"{{{_W}}}trPr"), f"{{{_W}}}gridBefore"
                ).set(f"{{{_W}}}val", "10000000")
            else:
                etree.SubElement(properties, f"{{{_W}}}vMerge")
            parts["word/document.xml"] = etree.tostring(root)

    _rewrite(path, mutate)
    with pytest.raises(ValueError):
        _parse(path)


@pytest.mark.parametrize("data", [b"", b"not a zip", b"PK\x03\x04truncated"])
def test_non_docx_bytes_are_errors(tmp_path: Path, data: bytes) -> None:
    path = tmp_path / "fake.docx"
    path.write_bytes(data)
    with pytest.raises(ValueError, match="DOCX"):
        _parse(path)


@pytest.mark.parametrize("suffix", ["xml", "dat"])
@pytest.mark.parametrize("budget", ["text", "elements"])
def test_docx_xml_budgets_follow_content_types_not_suffix(
    tmp_path: Path, suffix: str, budget: str
) -> None:
    from comet_rag.documents.parsers.docx.package import CT, RELS

    path = tmp_path / "renamed.docx"
    name = f"word/document.{suffix}"
    text = "x" * 1000 if budget == "text" else "x"
    paragraphs = 1 if budget == "text" else 100
    with ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            f'<Types xmlns="{CT}"><Override PartName="/{name}" '
            'ContentType="application/vnd.openxmlformats-officedocument.'
            'wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            f'<Relationships xmlns="{RELS}"><Relationship Id="r1" '
            f'Type="{_R}/officeDocument" Target="{name}"/></Relationships>',
        )
        archive.writestr(
            name,
            f'<w:document xmlns:w="{_W}"><w:body>'
            + f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>" * paragraphs
            + "</w:body></w:document>",
        )
    assert _parse(path).report.status == "complete"
    limits = (
        ArchiveLimits(max_xml_text_chars=64)
        if budget == "text"
        else ArchiveLimits(max_xml_elements=10)
    )
    with pytest.raises(ArchiveResourceLimitExceeded):
        _parse(path, archive_limits=limits)


def test_renamed_xml_part_uses_shared_budget_and_is_checked_once(
    tmp_path: Path,
) -> None:
    from comet_rag.documents.parsers.docx.package import CT, RELS, Package

    path = tmp_path / "aggregate.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr(
            "[Content_Types].xml",
            f'<Types xmlns="{CT}"><Override PartName="/word/document.dat" '
            'ContentType="application/vnd.openxmlformats-officedocument.'
            'wordprocessingml.document.main+xml"/></Types>',
        )
        archive.writestr(
            "_rels/.rels",
            f'<Relationships xmlns="{RELS}"><Relationship Id="r1" '
            f'Type="{_R}/officeDocument" Target="word/document.dat"/></Relationships>',
        )
        archive.writestr(
            "word/document.dat",
            f'<w:document xmlns:w="{_W}"><w:body><w:p><w:r><w:t>x</w:t>'
            "</w:r></w:p></w:body></w:document>",
        )
    with ZipFile(path) as archive:
        package = Package(archive, tmp_path, ArchiveLimits(max_xml_elements=9))
        root = package.xml(package.main)
        assert package.xml(package.main) is root
    with ZipFile(path) as archive:
        package = Package(archive, tmp_path, ArchiveLimits(max_xml_elements=8))
        with pytest.raises(ArchiveResourceLimitExceeded, match="elements"):
            package.xml(package.main)


def test_docx_budgets_reject_before_dom_and_materialization(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = build_image(tmp_path / "budget.docx")
    calls: list[object] = []
    original = etree.fromstring

    def fromstring(
        text: bytes, parser: etree.XMLParser | None = None
    ) -> etree._Element:
        calls.append(text)
        return original(text, parser=parser)

    monkeypatch.setattr(etree, "fromstring", fromstring)
    with pytest.raises(ArchiveResourceLimitExceeded):
        _parse(path, archive_limits=ArchiveLimits(max_members=1))
    assert calls == []
    assert not (tmp_path / "assets").exists()
    with pytest.raises(ValueError, match="exceed"):
        _parse(path, max_bytes=1)
    table = build_table(tmp_path / "wide.docx")
    with pytest.raises(ValueError, match="layout cells"):
        _parse(table, max_table_cells=8)


def test_field_hyperlinks_content_controls_and_native_list_levels(
    tmp_path: Path,
) -> None:
    doc = Document()
    paragraph = doc.add_paragraph("before ")
    for boundary, instruction, text in (
        ("begin", None, None),
        (None, ' HYPERLINK "https://example.org/field" ', None),
        ("separate", None, None),
        (None, None, "field label"),
        ("end", None, None),
    ):
        run = OxmlElement("w:r")
        if boundary:
            child = OxmlElement("w:fldChar")
            child.set(qn("w:fldCharType"), boundary)
        else:
            child = OxmlElement("w:instrText" if instruction else "w:t")
            child.text = instruction or text
        run.append(child)
        paragraph._p.append(run)
    control = OxmlElement("w:sdt")
    content = OxmlElement("w:sdtContent")
    control_run = OxmlElement("w:r")
    control_text = OxmlElement("w:t")
    control_text.text = " controlled "
    control_run.append(control_text)
    content.append(control_run)
    control.append(content)
    paragraph._p.append(control)
    doc.add_paragraph("one", style="List Number")
    nested = doc.add_paragraph("nested", style="List Number")
    numpr = OxmlElement("w:numPr")
    for name, value in (("numId", "5"), ("ilvl", "1")):
        node = OxmlElement(f"w:{name}")
        node.set(qn("w:val"), value)
        numpr.append(node)
    nested._p.get_or_add_pPr().append(numpr)
    path = tmp_path / "fields.docx"
    doc.save(str(path))

    def add_numbering_level(parts: dict[str, bytes]) -> None:
        root = etree.fromstring(parts["word/numbering.xml"])
        num = root.find(f"{{{_W}}}num[@{{{_W}}}numId='5']")
        assert num is not None
        reference = num.find(f"{{{_W}}}abstractNumId")
        assert reference is not None
        abstract = root.find(
            f"{{{_W}}}abstractNum[@{{{_W}}}abstractNumId='{reference.get(f'{{{_W}}}val')}']"
        )
        assert abstract is not None
        level = etree.SubElement(abstract, f"{{{_W}}}lvl", {f"{{{_W}}}ilvl": "1"})
        for name, value in (
            ("start", "1"),
            ("numFmt", "lowerLetter"),
            ("lvlText", "%1.%2)"),
        ):
            etree.SubElement(level, f"{{{_W}}}{name}").set(f"{{{_W}}}val", value)
        parts["word/numbering.xml"] = etree.tostring(root)

    _rewrite(path, add_numbering_level)
    result = _parse(path)
    rich = result.document.blocks[0].rich_text
    assert rich is not None
    assert [
        (item.text, item.link) for item in rich.inlines if isinstance(item, TextRun)
    ] == [
        ("before ", None),
        ("field label", "https://example.org/field"),
        (" controlled ", None),
    ]
    assert [
        block.level for block in result.document.blocks if block.kind == "list_item"
    ] == [0, 1]
    assert [
        block.marker for block in result.document.blocks if block.kind == "list_item"
    ] == ["1.", "1.a)"]
    assert result.report.status == "complete"


@pytest.mark.parametrize("story", ["body", "header", "footer"])
@pytest.mark.parametrize("hyperlink", [False, True], ids=["toc", "hyperlink"])
def test_fields_span_paragraphs_in_each_story(
    tmp_path: Path, story: str, hyperlink: bool
) -> None:
    doc = Document()
    container = doc if story == "body" else getattr(doc.sections[0], story)
    paragraph = container.add_paragraph()
    for boundary, instruction in (
        ("begin", None),
        (None, ' HYPERLINK "https://example.org/field" ' if hyperlink else " TOC "),
        ("separate", None),
    ):
        run = OxmlElement("w:r")
        child = OxmlElement("w:fldChar" if boundary else "w:instrText")
        if boundary:
            child.set(qn("w:fldCharType"), boundary)
        else:
            child.text = instruction
        run.append(child)
        paragraph._p.append(run)
    paragraph.add_run("first entry")
    paragraph = container.add_paragraph("second entry")
    run = OxmlElement("w:r")
    boundary = OxmlElement("w:fldChar")
    boundary.set(qn("w:fldCharType"), "end")
    run.append(boundary)
    paragraph._p.append(run)
    path = tmp_path / "spanning.docx"
    doc.save(str(path))

    result = _parse(path)
    entries = [block for block in result.document.blocks if _text(block)]
    assert [_text(block) for block in entries] == ["first entry", "second entry"]
    assert [block.role for block in entries] == [None if story == "body" else story] * 2
    assert all(
        isinstance(inline, TextRun)
        and inline.link == ("https://example.org/field" if hyperlink else None)
        for block in entries
        if block.rich_text
        for inline in block.rich_text.inlines
    )
    assert result.report.status == ("complete" if hyperlink else "partial")
    assert [issue.code for issue in result.report.issues] == (
        [] if hyperlink else ["unsupported_field"]
    )


def test_unclosed_fields_are_isolated_between_stories(tmp_path: Path) -> None:
    doc = Document()
    for container, text in (
        (doc, "body"),
        (doc.sections[0].header, "header"),
        (doc.sections[0].first_page_header, "first header"),
    ):
        paragraph = container.add_paragraph(text)
        run = OxmlElement("w:r")
        boundary = OxmlElement("w:fldChar")
        boundary.set(qn("w:fldCharType"), "begin")
        run.append(boundary)
        paragraph._p.append(run)
    doc.sections[0].footer.add_paragraph("footer")
    path = tmp_path / "isolated.docx"
    doc.save(str(path))

    result = _parse(path)
    assert [_text(block) for block in result.document.blocks if _text(block)] == [
        "body",
        "header",
        "first header",
        "footer",
    ]
    assert result.report.status == "partial"
    issues = result.report.issues
    assert [issue.code for issue in issues] == ["unclosed_field"] * 3
    parts: set[str] = set()
    for issue in issues:
        element_id = issue.locations[0].element_id
        assert element_id is not None
        parts.add(element_id.split(":")[0])
    assert len(parts) == 3


@pytest.mark.parametrize("leading_text", ["first ", ""], ids=["middle", "start"])
def test_inline_images_do_not_duplicate_list_items(
    tmp_path: Path, leading_text: str
) -> None:
    with ZipFile(build_image(tmp_path / "source.docx")) as archive:
        png = archive.read("word/media/image1.png")
    doc = Document()
    paragraph = doc.add_paragraph(leading_text, style="List Number")
    paragraph.add_run().add_picture(io.BytesIO(png))
    paragraph.add_run(" tail")
    doc.add_paragraph("second", style="List Number")
    path = tmp_path / "list-image.docx"
    doc.save(str(path))

    result = _parse(path)
    blocks = result.document.blocks
    assert [block.kind for block in blocks] == [
        "list_item",
        "image",
        "paragraph",
        "list_item",
    ]
    assert [_text(block) for block in blocks] == [leading_text, "", " tail", "second"]
    assert [block.marker for block in blocks] == ["1.", None, None, "2."]
    assert [block.level for block in blocks] == [0, None, None, 0]
    assert blocks[0].locations == blocks[2].locations
    assert result.report.status == "complete"
    assert loads(dumps(result)) == result


def test_external_images_and_unsupported_fields_have_partial_issues(
    tmp_path: Path,
) -> None:
    path = build_image(tmp_path / "external.docx")

    def mutate(parts: dict[str, bytes]) -> None:
        rels = etree.fromstring(parts["word/_rels/document.xml.rels"])
        for relation in rels:
            if relation.get("Type") == f"{_R}/image":
                relation.set("Target", "https://example.org/image.png")
                relation.set("TargetMode", "External")
        parts["word/_rels/document.xml.rels"] = etree.tostring(rels)
        root = etree.fromstring(parts["word/document.xml"])
        paragraph = root.find(f"{{{_W}}}body/{{{_W}}}p")
        assert paragraph is not None
        field = etree.SubElement(
            paragraph, f"{{{_W}}}fldSimple", {f"{{{_W}}}instr": " PAGE "}
        )
        etree.SubElement(etree.SubElement(field, f"{{{_W}}}r"), f"{{{_W}}}t").text = "9"
        parts["word/document.xml"] = etree.tostring(root)

    _rewrite(path, mutate)
    result = _parse(path)
    assert result.report.status == "partial"
    assert {issue.code for issue in result.report.issues} == {
        "unsupported_field",
        "external_image",
    }
    assert not result.document.assets
    assert _text(result.document.blocks[0]).endswith("9")


@pytest.mark.parametrize(
    "element_name", ["object", "sym", "footnoteReference", "ins", "del", "drawing"]
)
def test_known_unsupported_meaningful_elements_are_never_complete(
    tmp_path: Path, element_name: str
) -> None:
    path = build_basic(tmp_path / "unsupported.docx")

    def mutate(root: etree._Element) -> None:
        paragraph = root.find(f"{{{_W}}}body/{{{_W}}}p")
        assert paragraph is not None
        wrapper = etree.SubElement(paragraph, f"{{{_W}}}{element_name}")
        if element_name in {"ins", "del"}:
            etree.SubElement(
                etree.SubElement(wrapper, f"{{{_W}}}r"), f"{{{_W}}}t"
            ).text = "tracked text"
        elif element_name == "sym":
            wrapper.set(f"{{{_W}}}font", "Wingdings")
            wrapper.set(f"{{{_W}}}char", "F0B7")
        elif element_name == "footnoteReference":
            wrapper.set(f"{{{_W}}}id", "1")

    _mutate_document(path, mutate)
    result = _parse(path)
    assert result.report.status == "partial"
    assert result.report.issues
    assert all(issue.locations for issue in result.report.issues)
    if element_name in {"ins", "del"}:
        assert "tracked text" in project_document(result.document).text


def test_asset_conflicts_do_not_overwrite_existing_bytes(tmp_path: Path) -> None:
    path = build_image(tmp_path / "asset.docx")
    first = _parse(path)
    asset = first.document.assets[0]
    assert isinstance(asset.media.path, Path)
    asset.media.path.write_bytes(b"keep existing")
    with pytest.raises(ValueError, match="conflicts"):
        _parse(path)
    assert asset.media.path.read_bytes() == b"keep existing"


def test_bad_xml_and_expansion_budgets_fail_before_assets(tmp_path: Path) -> None:
    path = build_image(tmp_path / "budget.docx")
    for limits in (
        ArchiveLimits(max_xml_elements=1),
        ArchiveLimits(max_compression_ratio=1),
        ArchiveLimits(max_member_uncompressed_bytes=10),
    ):
        with pytest.raises(ArchiveResourceLimitExceeded):
            _parse(path, archive_limits=limits)
    assert not (tmp_path / "assets").exists()


def test_source_and_parser_cancellation_drain_blocking_work(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    path = tmp_path / "input.txt"
    path.write_bytes(b"input")
    resource = _resolve(path)

    async def check_parser() -> None:
        started = asyncio.Event()
        release = threading.Event()
        finished = threading.Event()
        loop = asyncio.get_running_loop()
        parser = AsyncTextParser()
        parse = parser._parse

        def blocked(item: LocalResource) -> ParseResult:
            loop.call_soon_threadsafe(started.set)
            assert release.wait(3)
            result = parse(item)
            finished.set()
            return result

        monkeypatch.setattr(parser, "_parse", blocked)
        task = asyncio.create_task(parser.parse(resource, output_dir=tmp_path))
        await asyncio.wait_for(started.wait(), 3)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert finished.is_set()
        assert path.read_bytes() == b"input"

    asyncio.run(check_parser())

    async def check_source() -> None:
        started = asyncio.Event()
        release = threading.Event()
        finished = threading.Event()
        loop = asyncio.get_running_loop()
        resolver = AsyncLocalSourceResolver()
        resolve = resolver._resolve

        def blocked(source: LocalSource) -> LocalResource:
            loop.call_soon_threadsafe(started.set)
            assert release.wait(3)
            result = resolve(source)
            finished.set()
            return result

        monkeypatch.setattr(resolver, "_resolve", blocked)

        async def enter() -> None:
            async with resolver.resolve(LocalSource(path)):
                pytest.fail("cancelled source must not be delivered")

        task = asyncio.create_task(enter())
        await asyncio.wait_for(started.wait(), 3)
        task.cancel()
        await asyncio.sleep(0)
        assert not task.done()
        release.set()
        with pytest.raises(asyncio.CancelledError):
            await task
        assert finished.is_set()
        assert path.exists()

    asyncio.run(check_source())


def test_invalid_public_source_and_parser_limits(tmp_path: Path) -> None:
    for value in (0, -1, True):
        with pytest.raises(ValueError):
            AsyncTextParser(max_bytes=value)
        with pytest.raises(ValueError):
            AsyncLocalSourceResolver(max_bytes=value)
        with pytest.raises(ValueError):
            AsyncDocxParser(max_table_cells=value)
    with pytest.raises(LookupError):
        AsyncTextParser(encoding="missing-encoding")
    with pytest.raises(ValueError):
        LocalSource("", source_id="id")
    with pytest.raises(ValueError):
        LocalSource(tmp_path, source_id=" ")


def test_empty_docx_is_complete_and_stable(tmp_path: Path) -> None:
    path = tmp_path / "blank.docx"
    Document().save(str(path))
    result = _parse(path)
    assert result.document.blocks == ()
    assert result.report.status == "complete"
    assert _parse(path).document.document_id == result.document.document_id
