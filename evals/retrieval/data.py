from __future__ import annotations

import io
import json
import zipfile
from dataclasses import dataclass
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from typing import Annotated, Literal

from docx import Document
from pydantic import BaseModel, ConfigDict, Field, model_validator

from comet_rag.engines.documents.docx.extractor import DocxDocumentExtractor
from comet_rag.engines.documents.normalization.markdown import (
    MarkdownDocumentNormalizer,
)
from comet_rag.ports.document import ExtractedDocument, NormalizedDocument

DATA = Path(__file__).parent
Name = Annotated[str, Field(pattern=r"^[a-z][a-z0-9_-]{0,63}$")]
Text = Annotated[str, Field(min_length=1)]


class StrictModel(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True, allow_inf_nan=False)


class Section(StrictModel):
    heading: Text
    paragraphs: list[Text]
    table: list[list[Text]] = Field(default_factory=list)


class Source(StrictModel):
    id: Name
    kb_id: Name
    source_id: Name
    revision: Name
    active: bool
    format: Literal["md", "docx", "pdf"]
    title: Text
    sections: list[Section]

    @model_validator(mode="after")
    def check_sections(self) -> Source:
        if self.format == "pdf" and self.sections:
            raise ValueError("PDF 只能使用外部提取快照，不能用 sections 伪造")
        if self.format != "pdf" and not self.sections:
            raise ValueError("生成式文档至少需要一个 section")
        return self


class Corpus(StrictModel):
    version: Literal[1]
    documents: Annotated[list[Source], Field(min_length=1)]


class Evidence(StrictModel):
    id: Name
    document_id: Name
    quote: Text


class Query(StrictModel):
    id: Name
    kb_id: Name
    text: Text
    category: Literal["fact", "paraphrase", "table", "multi", "no_answer"]
    relevance: dict[Name, Annotated[int, Field(ge=0, le=2)]]
    evidence: list[Name]


class Judgments(StrictModel):
    version: Literal[1]
    evidence: list[Evidence]
    queries: Annotated[list[Query], Field(min_length=1)]


def _unique_pairs(pairs: list[tuple[str, object]]) -> dict[str, object]:
    result: dict[str, object] = {}
    for key, value in pairs:
        if key in result:
            raise ValueError(f"重复 JSON 键：{key}")
        result[key] = value
    return result


def read_json[M: BaseModel](path: Path, model: type[M]) -> M:
    return model.model_validate(
        json.loads(path.read_text(encoding="utf-8"), object_pairs_hook=_unique_pairs)
    )


def write_json(path: Path, value: object) -> None:
    path.write_text(
        json.dumps(value, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def _docx(source: Source) -> bytes:
    doc = Document()
    doc.core_properties.created = datetime(2000, 1, 1)
    doc.core_properties.modified = datetime(2000, 1, 1)
    doc.add_heading(source.title, 1)
    for section in source.sections:
        doc.add_heading(section.heading, 2)
        for paragraph in section.paragraphs:
            doc.add_paragraph(paragraph)
        if section.table:
            table = doc.add_table(rows=0, cols=len(section.table[0]))
            for row in section.table:
                for cell, text in zip(table.add_row().cells, row, strict=True):
                    cell.text = text
    original = io.BytesIO()
    doc.save(original)
    canonical = io.BytesIO()
    # ZIP 时间戳也是生成物的一部分，不能让同一固定语料随生成时间变化。
    with zipfile.ZipFile(original) as src, zipfile.ZipFile(canonical, "w") as dst:
        for name in sorted(src.namelist()):
            item = zipfile.ZipInfo(name, date_time=(2000, 1, 1, 0, 0, 0))
            item.compress_type = zipfile.ZIP_DEFLATED
            item.external_attr = 0o600 << 16
            dst.writestr(item, src.read(name))
    return canonical.getvalue()


def _markdown(source: Source) -> str:
    blocks = [f"# {source.title}"]
    for section in source.sections:
        blocks.extend([f"## {section.heading}", *section.paragraphs])
        if section.table:
            header, *rows = section.table
            blocks.append(
                "\n".join(
                    "| " + " | ".join(row) + " |"
                    for row in [header, ["---"] * len(header), *rows]
                )
            )
    return "\n\n".join(blocks) + "\n"


@dataclass(frozen=True, slots=True)
class Dataset:
    corpus: Corpus
    judgments: Judgments
    documents: dict[str, NormalizedDocument]
    spans: dict[str, tuple[int, int]]
    fingerprint: str


def validate_sources(corpus: Corpus) -> dict[str, Source]:
    sources = {source.id: source for source in corpus.documents}
    if len(sources) != len(corpus.documents):
        raise ValueError("重复 document id")
    identities = {(s.kb_id, s.source_id, s.revision) for s in sources.values()}
    if len(identities) != len(sources):
        raise ValueError("重复 revision identity")
    if len({s.revision for s in sources.values()}) != len(sources):
        raise ValueError("revision id 必须全局唯一，符合 DocumentStore 契约")
    for identity in {(s.kb_id, s.source_id) for s in sources.values()}:
        if (
            sum(
                s.active for s in sources.values() if (s.kb_id, s.source_id) == identity
            )
            != 1
        ):
            raise ValueError("每个 source 必须恰有一个 active revision")
    for source in sources.values():
        for section in source.sections:
            if section.table and (
                not section.table[0]
                or any(len(row) != len(section.table[0]) for row in section.table)
            ):
                raise ValueError("表格列数不一致")
    return sources


def quote_locations(text: str, quote: str) -> tuple[int, ...]:
    """最多取两个匹配，足以区分缺失/唯一/歧义；重叠出现也算多次。"""
    if not quote:
        raise ValueError("证据 quote 不能为空")
    first = text.find(quote)
    if first < 0:
        return ()
    second = text.find(quote, first + 1)
    return (first,) if second < 0 else (first, second)


def locate_evidence(
    corpus: Corpus, labels: Judgments, documents: dict[str, NormalizedDocument]
) -> dict[str, tuple[int, int]]:
    sources = validate_sources(corpus)
    spans: dict[str, tuple[int, int]] = {}
    evidence = {item.id: item for item in labels.evidence}
    if len(evidence) != len(labels.evidence):
        raise ValueError("重复 evidence id")
    for item in labels.evidence:
        if item.document_id not in sources or not sources[item.document_id].active:
            raise ValueError("证据必须引用已知 active 文档")
        text = documents[item.document_id].markdown
        locations = quote_locations(text, item.quote)
        if len(locations) != 1:
            raise ValueError(f"证据必须唯一匹配规范正文：{item.id}")
        start = locations[0]
        spans[item.id] = (start, start + len(item.quote))
    if len({q.id for q in labels.queries}) != len(labels.queries):
        raise ValueError("重复 query id")
    for query in labels.queries:
        positive = {did for did, grade in query.relevance.items() if grade > 0}
        if (query.category == "no_answer") != (not positive):
            raise ValueError("no_answer 与正标签冲突")
        if bool(positive) != bool(query.evidence) or len(set(query.evidence)) != len(
            query.evidence
        ):
            raise ValueError("正标签需要非重复证据；无答案不得有证据")
        for did in query.relevance:
            if (
                did not in sources
                or not sources[did].active
                or sources[did].kb_id != query.kb_id
            ):
                raise ValueError("标签引用未知、inactive 或跨 KB 文档")
        if not positive and query.kb_id not in {s.kb_id for s in sources.values()}:
            raise ValueError("未知 query KB")
        if any(
            eid not in evidence or evidence[eid].document_id not in positive
            for eid in query.evidence
        ):
            raise ValueError("query evidence 必须属于正相关文档")
        if positive != {evidence[eid].document_id for eid in query.evidence}:
            raise ValueError("每个正相关文档必须有对应证据")
    return spans


def generate(output: Path, *, data: Path = DATA) -> Dataset:
    """只写新目录；标签先于切块定义，规范文本中必须能唯一定位证据。"""
    corpus = read_json(data / "corpus.json", Corpus)
    labels = read_json(data / "judgments.json", Judgments)
    sources = validate_sources(corpus)
    if any(source.format == "pdf" for source in sources.values()):
        raise ValueError(
            "真实 PDF 必须使用冻结页快照加载器，不能生成 Markdown 冒充 PDF"
        )
    # 原文件可用于失败诊断；只有解析和全部标签校验通过才发布 manifest 指纹。
    files: dict[str, bytes] = {}
    for source in sources.values():
        files[f"{source.id}.{source.format}"] = (
            _docx(source) if source.format == "docx" else _markdown(source).encode()
        )
    output.mkdir(parents=True, exist_ok=False)
    for name, content in files.items():
        (output / name).write_bytes(content)
    normalizer = MarkdownDocumentNormalizer()
    documents: dict[str, NormalizedDocument] = {}
    for source in sources.values():
        path = output / f"{source.id}.{source.format}"
        extracted = (
            DocxDocumentExtractor(include_images=False).extract(
                path,
                filename=path.name,
                media_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            if source.format == "docx"
            else ExtractedDocument(path.read_text(encoding="utf-8"))
        )
        documents[source.id] = normalizer.normalize(extracted)
    spans = locate_evidence(corpus, labels, documents)
    manifest = {
        "corpus": corpus.model_dump(),
        "judgments": labels.model_dump(),
        "normalized": {did: doc.markdown for did, doc in documents.items()},
        "files_sha256": {name: sha256(raw).hexdigest() for name, raw in files.items()},
        "evidence_spans": spans,
    }
    fingerprint = sha256(
        json.dumps(manifest, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()
    write_json(output / "manifest.json", {"sha256": fingerprint, **manifest})
    return Dataset(corpus, labels, documents, spans, fingerprint)
